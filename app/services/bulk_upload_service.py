"""File-upload-based bulk calling (Excel/CSV), as an alternative to Google Sheets.

Every function here is shaped as a drop-in match for the corresponding
google_sheets_service function (read_sheet_data / update_sheet_status /
update_sheet_feedback), so telephony_service.trigger_next_queued_call and
the /calls endpoints can dispatch between "a real Google Sheet URL" and
"an uploaded batch" purely by checking is_upload_reference(sheet_url) and
picking the matching pair of functions — no other call site needs to know
the storage differs.

An uploaded batch has no live external document, so status write-backs go
to the bulk_upload_rows table instead of a spreadsheet; the same
/calls/sheet-status endpoint already built for Google Sheets reads it back
for the campaign-progress UI.
"""

from __future__ import annotations

import csv
import io
import uuid
from typing import Any

from app.services.google_sheets_service import find_header_row

UPLOAD_PREFIX = "upload:"

_ROW_FIELDS = (
    "name", "phone", "language", "status",
    "investigated_course", "internship_or_course", "course_duration",
)


def is_upload_reference(sheet_url: str) -> bool:
    return (sheet_url or "").startswith(UPLOAD_PREFIX)


def _upload_id(sheet_url: str) -> str:
    return sheet_url[len(UPLOAD_PREFIX):]


def _parse_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    return list(csv.reader(text.splitlines()))


def _parse_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    records = []
    for row in sheet.iter_rows(values_only=True):
        records.append(["" if cell is None else str(cell).strip() for cell in row])
    workbook.close()
    return records


def parse_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Parse an uploaded .csv/.xlsx file into the same row shape read_sheet_data produces."""
    extension = (filename or "").rsplit(".", 1)[-1].lower()
    if extension == "csv":
        records = _parse_csv(content)
    elif extension in ("xlsx", "xlsm"):
        records = _parse_xlsx(content)
    else:
        raise ValueError("Unsupported file type; upload a .csv or .xlsx file")
    if not records:
        return []

    header_row_index, headers, header_map = find_header_row(records)
    results = []
    for idx, row in enumerate(records[header_row_index:], start=header_row_index + 1):
        if len(row) < len(headers):
            row = row + [""] * (len(headers) - len(row))
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        normalized: dict[str, Any] = {"row_number": idx}
        for key in _ROW_FIELDS:
            if key in header_map:
                normalized[key] = (row_dict.get(header_map[key]) or "").strip()
        results.append(normalized)
    return results


def create_batch(
    db, organization_id: int, campaign_id: int, filename: str, rows: list[dict[str, Any]],
) -> str:
    """Persist parsed rows and return a sheet_url-shaped reference for them."""
    upload_id = str(uuid.uuid4())
    cursor = db.cursor()
    for row in rows:
        phone = (row.get("phone") or "").strip()
        if not phone:
            continue
        cursor.execute(
            """INSERT INTO bulk_upload_rows
                   (upload_id, organization_id, campaign_id, row_index, name, phone, language,
                    investigated_course, internship_or_course, course_duration, status, source_filename)
               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'',%s)""",
            (
                upload_id, organization_id, campaign_id, row.get("row_number"),
                row.get("name"), phone, row.get("language"),
                row.get("investigated_course"), row.get("internship_or_course"),
                row.get("course_duration"), filename,
            ),
        )
    db.commit()
    return f"{UPLOAD_PREFIX}{upload_id}"


def read_rows(sheet_url: str) -> list[dict[str, Any]]:
    """Drop-in match for google_sheets_service.read_sheet_data's signature/shape."""
    from app.database.db import get_db_connection

    db = get_db_connection()
    try:
        cursor = db.cursor()
        cursor.execute(
            """SELECT row_index, name, phone, language, status,
                      investigated_course, internship_or_course, course_duration
               FROM bulk_upload_rows WHERE upload_id=%s ORDER BY row_index""",
            (_upload_id(sheet_url),),
        )
        rows = cursor.fetchall()
    finally:
        db.close()
    return [
        {
            "row_number": row["row_index"],
            "name": row.get("name") or "",
            "phone": row.get("phone") or "",
            "language": row.get("language") or "",
            "status": row.get("status") or "",
            "investigated_course": row.get("investigated_course") or "",
            "internship_or_course": row.get("internship_or_course") or "",
            "course_duration": row.get("course_duration") or "",
        }
        for row in rows
    ]


def update_status(sheet_url: str, row_number: int, status: str) -> None:
    """Drop-in match for google_sheets_service.update_sheet_status."""
    from app.database.db import get_db_connection

    db = get_db_connection()
    try:
        cursor = db.cursor()
        cursor.execute(
            "UPDATE bulk_upload_rows SET status=%s WHERE upload_id=%s AND row_index=%s",
            (status, _upload_id(sheet_url), row_number),
        )
        db.commit()
    finally:
        db.close()


def update_feedback(sheet_url: str, row_number: int, feedback_data: dict[str, Any]) -> None:
    """Drop-in match for google_sheets_service.update_sheet_feedback.

    There is no live document to write summary/transcript/etc. back to, so
    only status is tracked here — the rest already persists to call_logs.
    """
    status = feedback_data.get("status")
    if status:
        update_status(sheet_url, row_number, status)
